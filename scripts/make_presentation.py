"""
VIDYA-X AI — PowerPoint Generator Script
Generates presentation/VIDYA_X_AI_Hackathon_Pitch.pptx cleanly.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

DARK_BG = RGBColor(15, 23, 42)
CARD_BG = RGBColor(30, 41, 59)
ACCENT_BLUE = RGBColor(56, 189, 248)
TEXT_WHITE = RGBColor(248, 250, 252)
TEXT_MUTED = RGBColor(148, 163, 184)
GREEN_PASS = RGBColor(34, 197, 94)

def add_header(slide, title_text, category_text="VIDYA-X AI — HACKATHON PITCH"):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = category_text.upper()
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE
    
    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE

# Slide 1: Title & Vision
s1 = prs.slides.add_slide(blank_layout)
tb = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.5))
tf = tb.text_frame
p = tf.paragraphs[0]
p.text = "VIDYA-X AI"
p.font.size = Pt(48)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p2 = tf.add_paragraph()
p2.text = "Your Course-Aware, Adaptive & Multimodal AI Tutor"
p2.font.size = Pt(24)
p2.font.bold = True
p2.font.color.rgb = TEXT_WHITE

p3 = tf.add_paragraph()
p3.text = "\"Learn from your material. Get cited answers. Practice intelligently. Improve continuously.\""
p3.font.size = Pt(18)
p3.font.italic = True
p3.font.color.rgb = TEXT_MUTED

p4 = tf.add_paragraph()
p4.text = "\nMultimodal AI Hackathon 2026 | Track D: Personalized Tutoring & Adaptive Learning"
p4.font.size = Pt(14)
p4.font.color.rgb = ACCENT_BLUE

# Slide 2: The Problem
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "The Problem: Why Generic AI Fails University Students")
tb = s2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf = tb.text_frame
tf.word_wrap = True
bullets = [
    ("1. Ungrounded Hallucinations", "Generic LLMs fabricate formulas and definitions outside the student's syllabus."),
    ("2. Broken Source Traceability", "Students cannot verify which textbook page, slide, or video timestamp an explanation came from."),
    ("3. Static, One-Size-Fits-All Quizzes", "Standard quiz tools don't adapt difficulty or track cognitive skill acquisition."),
    ("4. Superficial Error Feedback", "Marking an answer 'incorrect' fails to diagnose underlying conceptual misconceptions."),
    ("5. Disconnected Study Habits", "Notes, flashcards, quizzes, and calendars exist in disconnected silos without spaced repetition.")
]
for title, desc in bullets:
    p = tf.add_paragraph()
    p.text = f"• {title}: "
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = ACCENT_BLUE
    run = p.add_run()
    run.text = desc
    run.font.bold = False
    run.font.color.rgb = TEXT_WHITE

# Slide 3: The Solution
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "The Solution: A Closed-Loop Cognitive Mastery Architecture")
tb = s3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf = tb.text_frame
tf.word_wrap = True
loop_steps = [
    "1. Multimodal Ingestion: Ingest PDFs, PPTX slides, and MP4 videos with coordinate preservation.",
    "2. Knowledge Structuring: Build course-isolated knowledge graphs with pgvector semantic embeddings.",
    "3. Grounded AI Tutor: Answer student questions with exact citations & safe out-of-domain abstention.",
    "4. Adaptive Practice: Generate MCQ, short-answer, and numerical quizzes matched to student mastery.",
    "5. BKT Learner Modeling: Track fine-grained latent concept mastery via pure Bayesian mathematics.",
    "6. Misconception Remediation: Isolate systematic errors with contrast cards, worked examples, & self-checks.",
    "7. Spaced Repetition Planner: Schedule daily Leitner flashcard reviews and exam-ready study tasks."
]
for step in loop_steps:
    p = tf.add_paragraph()
    p.text = f"• {step}"
    p.font.size = Pt(15)
    p.font.color.rgb = TEXT_WHITE

# Slide 4: Architecture
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "Technical Architecture: Production-Ready & Modular")
tb = s4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf = tb.text_frame
tf.word_wrap = True
arch_items = [
    ("Frontend Tier", "Zero-framework HTML5 / CSS3 / Vanilla ES-Module JavaScript (sub-millisecond loads)."),
    ("Application Tier", "FastAPI with Python 3.12, strict Pydantic validation, async endpoints, and health probes."),
    ("Persistence & Vector Store", "PostgreSQL 16 with pgvector HNSW indexing (with pure Python/SQLite fallback)."),
    ("Multimodal Extraction", "PyMuPDF (PDF), python-pptx (Slides), FFmpeg + faster-whisper (Video transcription)."),
    ("Cognitive Engines", "Pure mathematical Bayesian Knowledge Tracing (BKT) + Leitner spaced repetition.")
]
for title, desc in arch_items:
    p = tf.add_paragraph()
    p.text = f"• {title}: "
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = ACCENT_BLUE
    run = p.add_run()
    run.text = desc
    run.font.bold = False
    run.font.color.rgb = TEXT_WHITE

# Slide 5: Evaluation
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Evaluation & Verification: Honest, Reproducible Benchmarks")
tb = s5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf = tb.text_frame
tf.word_wrap = True
eval_items = [
    ("Automated Test Suite", "68 / 68 Tests Passing (100%) across BKT, Quiz, Adaptive, Remediation, Study Plan, and Visuals."),
    ("Course Isolation Release Gate", "0.0 Cross-Course Leakage Violations (Target: 0.0) — Strict SQL tenant filtering."),
    ("Unsupported Question Abstention", "1.0 (100%) Abstention Rate on out-of-domain queries — Verified anti-hallucination protection."),
    ("Reproducible Test Harness", "Versioned JSONL test cases in backend/evaluation/ with automated reporting."),
    ("Honest Engineering Disclosure", "Unit tests verify software correctness; longitudinal learning gains require human clinical trials.")
]
for title, desc in eval_items:
    p = tf.add_paragraph()
    p.text = f"• {title}: "
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = GREEN_PASS
    run = p.add_run()
    run.text = desc
    run.font.bold = False
    run.font.color.rgb = TEXT_WHITE

# Slide 6: Summary & Roadmap
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "Impact & Next Steps: From Hackathon to University Deployments")
tb = s6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
tf = tb.text_frame
tf.word_wrap = True
summary_items = [
    ("Production Hardened", "Multi-stage Docker, health probes (/health, /health/ready), and smoke test automation."),
    ("LMS Integration (Next)", "Connect via LTI 1.3 to university Canvas and Blackboard course shells."),
    ("Collaborative Peer Study", "Multi-student shared review decks and cohort-level misconception analytics for professors."),
    ("Core Philosophy", "\"Learn from your material. Get cited answers. Practice intelligently. Improve continuously.\"")
]
for title, desc in summary_items:
    p = tf.add_paragraph()
    p.text = f"• {title}: "
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = ACCENT_BLUE
    run = p.add_run()
    run.text = desc
    run.font.bold = False
    run.font.color.rgb = TEXT_WHITE

os.makedirs("presentation", exist_ok=True)
output_path = "presentation/VIDYA_X_AI_Hackathon_Pitch.pptx"
prs.save(output_path)
print(f"Successfully saved {output_path}")
